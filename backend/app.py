"""
SmartRoute Backend — Streamlit mini-app for testing the pipeline.

Run:  streamlit run backend/app.py

Accepts raw text, Excel (.xlsx), or images (.png/.jpg) as input,
runs the full pipeline (extractor → validator → urgency), and
displays the resulting InspectionRecord as JSON.
"""

import json
import sys
from pathlib import Path

# Add project root to path so `backend.*` imports work
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import streamlit as st
import openpyxl
import pytesseract
from PIL import Image

from backend.extractor import extract_fields
from backend.validator import validate
from backend.urgency import apply_routing


# ─── File Parsers ────────────────────────────────────────────────────────────

def parse_xlsx(uploaded_file) -> str:
    """Read an Excel file and concatenate all cell values into a single string."""
    wb = openpyxl.load_workbook(uploaded_file, read_only=True)
    lines = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" | ".join(cells))
    wb.close()
    return "\n".join(lines)


def parse_image(uploaded_file) -> str:
    """Run OCR on an uploaded image and return the extracted text."""
    image = Image.open(uploaded_file)
    return pytesseract.image_to_string(image)


# ─── Pipeline Runner ─────────────────────────────────────────────────────────

def run_pipeline(raw_text: str) -> dict:
    """Run the full SmartRoute pipeline and return the record as a dict."""
    extraction = extract_fields(raw_text)
    record = validate(extraction, raw_text)
    record = apply_routing(record)
    return record.model_dump(mode="json")


# ─── Streamlit UI ────────────────────────────────────────────────────────────

def main():
    st.set_page_config(page_title="SmartRoute Backend", layout="wide")
    st.title("SmartRoute Backend")
    st.caption("Test the extraction pipeline without the full frontend.")

    input_mode = st.radio("Input type", ["Raw Text", "Excel (.xlsx)", "Image (.png/.jpg)"])

    raw_text = None

    if input_mode == "Raw Text":
        raw_text = st.text_area("Paste inspection text", height=200)

    elif input_mode == "Excel (.xlsx)":
        uploaded = st.file_uploader("Upload Excel file", type=["xlsx"])
        if uploaded:
            try:
                raw_text = parse_xlsx(uploaded)
                st.subheader("Extracted text")
                st.text(raw_text)
            except Exception as e:
                st.error(f"Failed to read Excel file: {e}")

    elif input_mode == "Image (.png/.jpg)":
        uploaded = st.file_uploader("Upload image", type=["png", "jpg", "jpeg"])
        if uploaded:
            st.image(uploaded, width=400)
            try:
                raw_text = parse_image(uploaded)
                st.subheader("OCR text")
                st.text(raw_text)
            except Exception as e:
                st.error(f"OCR failed: {e}")

    if st.button("Process", disabled=not raw_text):
        if not raw_text or not raw_text.strip():
            st.warning("No text to process.")
            return

        with st.spinner("Running pipeline..."):
            try:
                result = run_pipeline(raw_text)
                st.subheader("InspectionRecord")
                st.json(json.dumps(result, indent=2, default=str))
            except ConnectionError as e:
                st.error(f"Ollama connection error: {e}")
            except ValueError as e:
                st.error(f"Extraction error: {e}")


if __name__ == "__main__":
    main()

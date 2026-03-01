import re
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image
import pytesseract


def safe_name(name: str) -> str:
    name = re.sub(r"[^\w\-. ]+", "_", name)
    return name.strip().replace(" ", "_")


def dump_xlsx_to_txt(xlsx_path: Path, out_dir: Path) -> None:
    wb = load_workbook(xlsx_path, data_only=True)
    out_path = out_dir / f"{safe_name(xlsx_path.stem)}.txt"

    with out_path.open("w", encoding="utf-8") as f:
        f.write(f"FILE: {xlsx_path.name}\n")
        f.write("=" * 80 + "\n\n")

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            f.write(f"[SHEET] {sheet_name}\n")
            f.write("-" * 80 + "\n")

            max_row = ws.max_row
            max_col = ws.max_column

            for r in range(1, max_row + 1):
                row_vals = []
                empty_row = True
                for c in range(1, max_col + 1):
                    v = ws.cell(row=r, column=c).value
                    if v is None:
                        row_vals.append("")
                    else:
                        empty_row = False
                        row_vals.append(str(v))
                if empty_row:
                    continue
                f.write("\t".join(row_vals) + "\n")

            f.write("\n\n")

    print(f"[XLSX] Wrote: {out_path}")


def dump_png_to_txt(png_path: Path, out_dir: Path) -> None:
    out_path = out_dir / f"{safe_name(png_path.stem)}.txt"

    img = Image.open(png_path).convert("RGB")
    text = pytesseract.image_to_string(img)

    with out_path.open("w", encoding="utf-8") as f:
        f.write(f"FILE: {png_path.name}\n")
        f.write("=" * 80 + "\n\n")
        f.write(text.strip() + "\n")

    print(f"[PNG/OCR] Wrote: {out_path}")


def main():
    import argparse

    parser = argparse.ArgumentParser(description="Extract XLSX + PNG content into TXT files.")
    parser.add_argument("input_dir", help="Folder containing .xlsx and .png files")
    parser.add_argument("--out", default="output_txt", help="Output folder (created if missing)")
    parser.add_argument(
        "--tesseract",
        default=r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        help="Full path to tesseract.exe",
    )

    args = parser.parse_args()

    in_dir = Path(args.input_dir).expanduser().resolve()
    out_dir = Path(args.out).expanduser().resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    
    # Set Tesseract path explicitly (Windows-safe)
    pytesseract.pytesseract.tesseract_cmd = args.tesseract

    if not in_dir.exists() or not in_dir.is_dir():
        raise SystemExit(f"Input directory not found: {in_dir}")

    xlsx_files = sorted(in_dir.glob("*.xlsx"))
    png_files = sorted(in_dir.glob("*.png"))

    if not xlsx_files and not png_files:
        print("No .xlsx or .png files found in:", in_dir)
        return

    for x in xlsx_files:
        try:
            dump_xlsx_to_txt(x, out_dir)
        except Exception as e:
            print(f"[XLSX] Failed {x.name}: {e}")

    for p in png_files:
        try:
            dump_png_to_txt(p, out_dir)
        except Exception as e:
            print(f"[PNG] Failed {p.name}: {e}")

    print("\nDone.")
    print("Output folder:", out_dir)


if __name__ == "__main__":
    main()